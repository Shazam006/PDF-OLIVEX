"""Runs against real installed engines, including inside the Docker test target."""
import io

import fitz
import pytest
from PIL import Image, ImageDraw, ImageFont
from backend.runtime import system_capabilities

caps=system_capabilities()["tools"]


@pytest.mark.skipif(not caps["ocr"],reason="Real OCR engines are not installed")
@pytest.mark.parametrize("language",["por","eng","spa"])
def test_real_ocr(client,language):
    image=Image.new("RGB",(1200,500),"white")
    draw=ImageDraw.Draw(image)
    font=ImageFont.load_default(size=60)
    draw.text((80,140),"PDF OLIVEX DOCUMENT TEST",font=font,fill="black")
    stream=io.BytesIO();image.save(stream,"PNG")
    response=client.post("/api/scan-to-pdf",files={"files":("scan.png",stream.getvalue(),"image/png")},data={"run_ocr":"true","language":language})
    assert response.status_code==200,response.text[:1000] if response.status_code!=200 else ""
    with fitz.open(stream=response.content,filetype="pdf") as doc:
        assert "OLIVEX" in doc[0].get_text().upper()


@pytest.mark.skipif(not caps["office"],reason="Real LibreOffice engine is not installed")
@pytest.mark.parametrize("kind",["docx","xlsx","pptx"])
def test_real_office_conversion(client,kind):
    stream=io.BytesIO()
    if kind=="docx":
        from docx import Document
        document=Document();document.add_paragraph("PDF OLIVEX OFFICE TEST");document.save(stream)
    elif kind=="xlsx":
        from openpyxl import Workbook
        document=Workbook();document.active["A1"]="PDF OLIVEX OFFICE TEST";document.save(stream)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document=Presentation();slide=document.slides.add_slide(document.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text="PDF OLIVEX OFFICE TEST";document.save(stream)
    response=client.post("/api/office-to-pdf",files={"file":("document."+kind,stream.getvalue(),"application/octet-stream")})
    assert response.status_code==200,response.text[:1000] if response.status_code!=200 else ""
    with fitz.open(stream=response.content,filetype="pdf") as doc:
        assert "OLIVEX" in " ".join(page.get_text() for page in doc)


@pytest.mark.skipif(not caps["pdfa"],reason="Real PDF/A engine is not installed")
def test_real_pdfa_conversion(client,pdf):
    response=client.post("/api/pdf-a",files={"file":("source.pdf",pdf,"application/pdf")})
    assert response.status_code==200,response.text[:1000] if response.status_code!=200 else ""
    with fitz.open(stream=response.content,filetype="pdf") as doc:
        assert "pdfaid:part" in doc.get_xml_metadata() and len(doc)==3
