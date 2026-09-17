import io
import os
import sys
from pathlib import Path

import fitz
from fastapi import UploadFile, File, Form, HTTPException

from .runtime import MAX_PAGES, download, open_pdf, out, require_tool, run_tool, save_upload, tool_path


def html_to_pdf(file: UploadFile = File(...)):
    from lxml import html, etree
    source = save_upload(file, {".html", ".htm"})
    try:
        tree = html.fromstring(Path(source).read_bytes())
    except (etree.ParserError, ValueError):
        raise HTTPException(400, "HTML inválido.")
    for node in tree.xpath("//script|//iframe|//object|//embed|//link|//img|//base"):
        node.drop_tree()
    markup = html.tostring(tree, encoding="unicode")
    path = out()
    paper = fitz.paper_rect("a4")
    writer = fitz.DocumentWriter(path)
    try:
        story = fitz.Story(markup)
        def layout(number, filled):
            if number >= MAX_PAGES:
                raise HTTPException(400, "HTML excede o limite de páginas.")
            return paper, paper + (36,36,-36,-36), None
        story.write(writer, layout)
    finally:
        writer.close()
    return download(path, "html_convertido.pdf")


def pdf_a(file: UploadFile = File(...)):
    from .main import execute_ocr
    source = save_upload(file)
    with open_pdf(source):
        pass
    path = execute_ocr(source, "por", "pdfa-2")
    with open_pdf(path) as doc:
        if "pdfaid:part" not in doc.get_xml_metadata():
            raise HTTPException(400, "A conversão não produziu metadados PDF/A.")
    validation = "not-independent"
    if validator := tool_path(["verapdf"]):
        result = run_tool([validator, "--format", "xml", path])
        from lxml import etree
        tree = etree.fromstring(result.stdout)
        reports = tree.xpath("//*[local-name()='validationReport']")
        if not reports or any(report.get("isCompliant") != "true" for report in reports):
            raise HTTPException(400, "O validador não confirmou a conformidade PDF/A.")
        validation = "passed"
    return download(path, "arquivo_pdfa.pdf", headers={"X-PDFA-Validation":validation})


def pdf_to_office(file: UploadFile = File(...), target: str = Form("docx")):
    if target not in {"docx", "xlsx", "pptx"}:
        raise HTTPException(400, "Formato Office inválido.")
    require_tool(target)
    source = save_upload(file)
    path = out("."+target)
    with open_pdf(source) as doc:
        if target == "docx":
            if not any(page.get_text().strip() for page in doc):
                raise HTTPException(400, "O PDF não tem texto pesquisável. Aplique OCR antes de converter para Word.")
            run_tool([sys.executable, str(Path(__file__).with_name("docx_worker.py")), source, path])
        elif target == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.remove(workbook.active)
            for number,page in enumerate(doc,1):
                for index,table in enumerate(page.find_tables().tables,1):
                    sheet = workbook.create_sheet(f"P{number}_Tabela{index}")
                    for row in table.extract():
                        sheet.append([("'"+value) if isinstance(value,str) and value.startswith(("=","+","-","@")) else value for value in row])
            if not workbook.sheetnames:
                raise HTTPException(400, "Nenhuma tabela detectada. PDFs escaneados precisam de OCR; tabelas sem estrutura podem exigir revisão.")
            workbook.save(path)
        else:
            from pptx import Presentation
            from pptx.util import Inches
            presentation = Presentation()
            presentation.slide_width, presentation.slide_height = Inches(13.33), Inches(7.5)
            for page in doc:
                scale = min(2, (8_000_000 / (page.rect.width*page.rect.height))**.5)
                pix = page.get_pixmap(matrix=fitz.Matrix(scale,scale), alpha=False)
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                ratio = min(presentation.slide_width/pix.width, presentation.slide_height/pix.height)
                width,height = int(pix.width*ratio),int(pix.height*ratio)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")),
                    int((presentation.slide_width-width)/2),int((presentation.slide_height-height)/2),width=width,height=height)
            presentation.save(path)
    media = {"docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             "xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
             "pptx":"application/vnd.openxmlformats-officedocument.presentationml.presentation"}[target]
    return download(path,"convertido."+target,media)


def sign(file: UploadFile = File(...), certificate: UploadFile = File(...), password: str = Form(""), reason: str = Form("")):
    require_tool("sign")
    source = save_upload(file)
    with open_pdf(source):
        pass
    cert = save_upload(certificate, {".p12", ".pfx"})
    from pyhanko.sign import signers
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    try:
        signer = signers.SimpleSigner.load_pkcs12(cert,passphrase=password.encode())
        if signer is None:
            raise ValueError
    except Exception:
        raise HTTPException(400, "Certificado inválido ou senha incorreta.")
    path = out()
    try:
        with open(source,"rb") as inp, open(path,"wb") as output:
            signers.sign_pdf(IncrementalPdfFileWriter(inp),
                signers.PdfSignatureMetadata(field_name="AssinaturaOLIVEX"+os.urandom(4).hex(), reason=reason[:200] or None),
                signer=signer,output=output)
    except Exception:
        raise HTTPException(400, "Não foi possível assinar este documento com o certificado informado.")
    return download(path,"pdf_assinado.pdf")
